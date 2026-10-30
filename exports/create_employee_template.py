from pathlib import Path
import re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

root = Path(__file__).resolve().parents[1]
def options(filename):
    return re.findall(r'\{ value: "([^"]+)", label: "([^"]+)" \}', (root/'frontend/src/components/ui'/filename).read_text(encoding='utf-8'))
headers = ['Mã nhân viên *','Họ và tên *','Chức vụ *','Phòng ban *','Cấp bậc *','Giới tính *','Ngày gia nhập *','Ngày sinh','Văn phòng / Địa điểm làm việc *','Trạng thái làm việc *','Email công việc *','Số điện thoại *','Ảnh đại diện (tên file hoặc URL)','Mô tả / Tiểu sử công việc']
choices = {4:options('employee-department-options.ts'),5:options('employee-rank-options.ts'),6:options('employee-gender-options.ts'),9:[('Văn Phòng Á Châu Dĩ An','Văn Phòng Á Châu Dĩ An'),('Bình Dương','Bình Dương (Nhà máy số 1)'),('Long An','Long An (Nhà máy số 2)')],10:options('employee-status-options.ts')}
w=Workbook();s=w.active;s.title='Nhập nhân viên';g=w.create_sheet('Hướng dẫn');d=w.create_sheet('Danh mục');e=w.create_sheet('Ví dụ')
s.append(headers);e.append(headers);d.append(['Trường','Lựa chọn hiển thị','Mã hệ thống'])
def validation(col,kind,formula,message,optional=False):
    v=DataValidation(type=kind,formula1=formula,allow_blank=optional,errorStyle='stop',showErrorMessage=True,showInputMessage=True,showDropDown=False)
    v.errorTitle='Thông tin chưa hợp lệ';v.error=message;v.promptTitle='Nhập thông tin';v.prompt=message
    s.add_data_validation(v);letter=get_column_letter(col);v.add(f'{letter}2:{letter}501');return v
for col,entries in choices.items():
    assert entries
    start=d.max_row+1
    for value,label in entries:d.append([headers[col-1].replace(' *',''),label,value])
    name=f'Options{col}';w.defined_names.add(DefinedName(name,attr_text=f"'Danh mục'!$B${start}:$B${d.max_row}"))
    validation(col,'list','='+name,'Chọn từ danh sách bằng mũi tên trong ô hoặc Alt + mũi tên xuống.')
    letter=get_column_letter(col)
    s.conditional_formatting.add(f'{letter}2:{letter}501',FormulaRule(formula=[f'AND({letter}2<>"",COUNTIF({name},{letter}2)=0)'],fill=PatternFill('solid',fgColor='FFC7CE')))
def digits(cell):
    for n in '0123456789':cell=f'SUBSTITUTE({cell},"{n}","")'
    return f'LEN({cell})=0'
rules={1:(f'AND(LEN(A2)>0,{digits("A2")},COUNTIF($A$2:$A$501,A2)=1)','Mã chỉ gồm chữ số, không trùng trong file.'),2:('AND(LEN(TRIM(B2))>=2,LEN(B2)<=100)','Họ tên từ 2 đến 100 ký tự.'),3:('AND(LEN(TRIM(C2))>=2,LEN(C2)<=100)','Chức vụ từ 2 đến 100 ký tự.'),11:('AND(LEN(K2)>12,RIGHT(LOWER(K2),12)="@asiafnb.com",LEN(K2)-LEN(SUBSTITUTE(K2,"@",""))=1,ISERROR(SEARCH(" ",K2)))','Email phải có đuôi @asiafnb.com, không chứa khoảng trắng.'),12:(f'AND(LEN(L2)=10,LEFT(L2,1)="0",ISNUMBER(FIND(MID(L2,2,1),"35789")),{digits("L2")})','Điện thoại gồm 10 số, đầu 03, 05, 07, 08 hoặc 09.'),14:('LEN(N2)<=2000','Mô tả tối đa 2.000 ký tự.')}
for col,(formula,message) in rules.items():
    validation(col,'custom',formula,message,col==14)
    letter=get_column_letter(col);s.conditional_formatting.add(f'{letter}2:{letter}501',FormulaRule(formula=[f'AND({letter}2<>"",NOT({formula}))'],fill=PatternFill('solid',fgColor='FFC7CE')))
for col in (7,8):
    letter=get_column_letter(col)
    cell=f'{letter}2'
    day=f'VALUE(LEFT({cell},2))'
    month=f'VALUE(MID({cell},4,2))'
    year=f'VALUE(RIGHT({cell},4))'
    date=f'DATE({year},{month},{day})'
    formula=f'IFERROR(AND(ISTEXT({cell}),LEN({cell})=10,MID({cell},3,1)="/",MID({cell},6,1)="/",{year}>=1900,{year}<=2100,TEXT({date},"dd/mm/yyyy")={cell}),FALSE)'
    validation(col,'custom',formula,'Nhập ngày dưới dạng DD/MM/YYYY, ví dụ 25/10/2005.',col==8)
    s.conditional_formatting.add(f'{letter}2:{letter}501',FormulaRule(formula=[f'AND({cell}<>"",NOT({formula}))'],fill=PatternFill('solid',fgColor='FFC7CE')))
for col in [1,2,3,4,5,6,7,9,10,11,12]:
    letter=get_column_letter(col);s.conditional_formatting.add(f'{letter}2:{letter}501',FormulaRule(formula=[f'AND(COUNTA($A2:$N2)>0,{letter}2="")'],fill=PatternFill('solid',fgColor='FFEB9C')))
s.conditional_formatting.add('E2:E501',FormulaRule(formula=['OR(AND(E2="CEO",D2<>"Phòng BOD"),AND(E2="CEO",COUNTIF($E$2:$E$501,"CEO")>1))'],fill=PatternFill('solid',fgColor='FFC7CE')))
for row in s.iter_rows(min_row=2,max_row=501,max_col=14):
    for c in row:
        c.number_format='@'
        if c.column in choices:c.fill=PatternFill('solid',fgColor='EAF4FC')
for row in [('HƯỚNG DẪN NHẬP NHÂN VIÊN','ASIA PORTAL'),('Cách dùng','Điền sheet Nhập nhân viên từ dòng 2. Mỗi dòng một người. Không sửa tên hoặc thứ tự cột. Tối đa 500 nhân viên.'),('Danh sách chọn','Chọn phòng ban, cấp bậc, giới tính, địa điểm và trạng thái từ danh sách trong ô.'),('Kiểm tra','Dấu * là bắt buộc. Mã, email và số điện thoại không được trùng nhau hoặc trùng hồ sơ hiện có, kể cả thùng rác. Nếu một dòng sai, toàn bộ file không được nhập.'),('Mã và điện thoại','Nhập dưới dạng văn bản để giữ số 0 đầu. Mã nhân viên chỉ gồm chữ số.'),('Ngày tháng','Nhập trực tiếp dưới dạng văn bản DD/MM/YYYY, ví dụ 25/10/2005 hoặc 01/02/2026. Không đổi định dạng ô sang Date. Ngày gia nhập bắt buộc; ngày sinh tùy chọn.'),('Cấp bậc','CEO phải thuộc Phòng BOD. Công ty chỉ có một CEO.'),('Ảnh','Để trống cột ảnh; cập nhật ảnh trong trang chỉnh sửa sau khi nhập.'),('Mô tả','Văn bản thuần, tối đa 2.000 ký tự.'),('Ví dụ','Sheet Ví dụ chỉ minh họa, không được nhập. Điền dữ liệu vào sheet Nhập nhân viên.')]:g.append(row)
e.append(['000123','Nguyễn Văn An','Chuyên viên nhân sự','Phòng Nhân sự & Hành chính','Chuyên viên/Nhân viên','Nam',None,None,'Văn Phòng Á Châu Dĩ An','Đang làm việc','nguyenvanan@asiafnb.com','0901234567',None,'Phụ trách tuyển dụng.'])
e['G2']='01/10/2026';e['G2'].number_format='@'
for sheet in w:
    for row in sheet:
        for c in row:c.font=Font(name='Arial',size=11);c.alignment=Alignment(wrap_text=True,vertical='top')
    for c in sheet[1]:c.font=Font(name='Arial',size=11,bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='087A43')
    sheet.row_dimensions[1].height=42;sheet.freeze_panes='A2'
for sheet in (s,e):
    for i,width in enumerate([20,28,30,25,29,16,22,22,38,26,36,24,40,56],1):sheet.column_dimensions[get_column_letter(i)].width=width
    sheet.freeze_panes='C2';sheet.sheet_view.zoomScale=85
g.column_dimensions['A'].width=28;g.column_dimensions['B'].width=110
for i in range(2,g.max_row+1):g.row_dimensions[i].height=60
for col,width in zip('ABC',[32,42,50]):d.column_dimensions[col].width=width
out=root/'frontend/public/Mau_nhap_nhan_vien_AsiaPortal.xlsx';w.save(out)
check=load_workbook(out);assert check.worksheets[0]['A1'].value=='Mã nhân viên *';assert len(check.defined_names)==5;assert len(check.worksheets[0].data_validations.dataValidation)==13
for col in ('G','H'):
    assert check.worksheets[0][f'{col}2'].number_format == '@'
    rule=next(v for v in check.worksheets[0].data_validations.dataValidation if f'{col}2' in v)
    assert rule.type == 'custom' and len(rule.formula1) <= 255
print('Created and validated: '+str(out))
